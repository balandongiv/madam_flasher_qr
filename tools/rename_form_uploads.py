"""Rename Microsoft Forms photo uploads using the Forms response Excel.

Copies (never renames/moves) each uploaded file from SRC to DST as
"<MATRIC>_<FULL NAME>.<ext>" and writes File_Name_Mapping.xlsx in DST.
Ends with an ATTENTION report of anything that needs a human to check.

Usage:
    python rename_form_uploads.py [SRC_FOLDER] [RESPONSES_XLSX] [DST_FOLDER]
Defaults are the 2025/2026 Majlis Anugerah Dekan paths below.
Requires: pip install openpyxl
"""
import os
import re
import shutil
import sys
from urllib.parse import unquote

import openpyxl
from openpyxl.styles import Font, PatternFill

SRC = r'C:/Users/balan/OneDrive - ums.edu.my/Question'
XLSX = SRC + '/MAJLIS ANUGERAH DEKAN & APRESIASI MAHASISWA SESSION 2025_2026(1-170).xlsx'
DST = r'C:/Users/balan/OneDrive/Desktop/new_name'

# Column headers in the Forms export are long; match them by prefix.
COLS = {
    'email': 'Email',
    'form_name': 'Name',
    'full_name': 'FULL NAME',
    'matric': 'MATRIC NUMBER',
    'programme': 'PROGRAMME',
    'photo': 'Please Upload a Photo',
}


def find_cols(header):
    idx = {}
    for key, prefix in COLS.items():
        for i, h in enumerate(header):
            if h and str(h).strip().startswith(prefix) and (key != 'form_name' or str(h).strip() == 'Name'):
                idx[key] = i
                break
        else:
            sys.exit(f'Column starting with "{prefix}" not found in responses Excel.')
    return idx


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else SRC
    xlsx = sys.argv[2] if len(sys.argv) > 2 else XLSX
    dst = sys.argv[3] if len(sys.argv) > 3 else DST
    os.makedirs(dst, exist_ok=True)

    ws = openpyxl.load_workbook(xlsx).active
    rows = list(ws.iter_rows(values_only=True))
    c = find_cols(rows[0])

    local = {f for f in os.listdir(src)
             if not f.lower().endswith(('.xlsx', '.csv')) and os.path.isfile(os.path.join(src, f))}
    used, seen, out = set(), set(), []
    attention = []

    for r in rows[1:]:
        if not r[c['photo']]:
            continue
        name = str(r[c['full_name']] or r[c['form_name']] or '').strip()
        matric = str(r[c['matric']] or '').strip().upper()
        if not matric:
            attention.append(f'No matric number typed by {name} ({r[c["email"]]}).')
        if matric and not re.fullmatch(r'[A-Z]{2}\d{8}', matric):
            attention.append(f'Unusual matric number "{matric}" for {name}.')
        for url in re.split(r'[;\s]+(?=https?://)', str(r[c['photo']]).strip()):
            orig = unquote(url.strip().rsplit('/', 1)[-1])
            ext = os.path.splitext(orig)[1].lower()
            base = re.sub(r'[\\/:*?"<>|]', '', f'{matric}_{name}'.strip('_'))
            new, k = base, 2
            while new + ext in seen:
                new = f'{base} ({k})'
                k += 1
            if orig in local:
                shutil.copy2(os.path.join(src, orig), os.path.join(dst, new + ext))
                used.add(orig)
                seen.add(new + ext)
                status = 'Copied'
                if new != base:
                    attention.append(f'{name} submitted more than once; saved as "{new + ext}".')
            else:
                status = 'Not found in folder'
                attention.append(f'{name} ({matric}): uploaded "{orig}" but it is not in the source folder '
                                 f'- download it from the link in the responses Excel.')
            out.append([orig, new + ext, r[c['form_name']], r[c['email']], name, matric,
                        r[c['programme']], status])

    extra = sorted(local - used)
    for f in extra:
        attention.append(f'"{f}" is in the source folder but matches no form response '
                         f'(likely an older/replaced upload) - not copied.')

    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'Mapping'
    s.append(['Original File Name', 'New File Name', 'Submitted By (Form Name)', 'Submitter Email',
              'Full Name (as typed)', 'Matric No.', 'Programme', 'Status'])
    for o in out:
        s.append(o)
    for f in extra:
        s.append([f, '', '', '', '', '', '', 'No matching form response'])
    warn = PatternFill('solid', fgColor='FFF2CC')
    for row in s.iter_rows(min_row=2):
        if row[7].value != 'Copied':
            for cell in row:
                cell.fill = warn
    for col, w in zip('ABCDEFGH', [55, 45, 35, 42, 35, 14, 30, 24]):
        s.column_dimensions[col].width = w
    for cell in s[1]:
        cell.font = Font(bold=True)
    s.freeze_panes = 'A2'
    s.auto_filter.ref = s.dimensions
    a = wb.create_sheet('Attention')
    a.append(['Items needing attention'])
    a['A1'].font = Font(bold=True)
    a.column_dimensions['A'].width = 140
    for m in attention:
        a.append([m])
    wb.save(os.path.join(dst, 'File_Name_Mapping.xlsx'))

    copied = sum(o[-1] == 'Copied' for o in out)
    print(f'Uploads in responses: {len(out)} | copied: {copied} | output: {dst}')
    print(f'\nATTENTION ({len(attention)}):' if attention else '\nATTENTION: none')
    for m in attention:
        print(' -', m)


if __name__ == '__main__':
    main()
